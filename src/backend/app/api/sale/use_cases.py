from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import re
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.exceptions.exceptions import (
    BadRequestException,
    FreemiumQuotaExceededException,
    InvalidImportDataException,
    ProductNotFoundException,
    SaleNotFoundException,
)
from app.api.sale.repository import SaleRepository
from app.api.sale.schema import (
    QuotaInfo,
    SaleCreateRequest,
    SaleCreateResponse,
    SaleHistoryResponse,
    SaleImportExcelResponse,
    SaleImportJsonRequest,
    SaleListResponse,
    SaleResponse,
)
from app.models.product_model import Product
from app.models.sale_model import Sale
from app.models.user_model import User


FREE_QUOTA = 300
ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

DATE_HEADER_ALIASES = {"tanggal", "date", "tgl"}
MENU_HEADER_ALIASES = {
    "menu",
    "nama menu",
    "nama_menu",
    "produk",
    "nama produk",
    "nama_produk",
    "product",
    "product_name",
}
QTY_HEADER_ALIASES = {
    "jumlah",
    "jumlah porsi terjual",
    "jumlah_porsi_terjual",
    "jumlah porsi",
    "porsi",
    "qty",
    "kuantitas",
}


def _to_sale_response(sale: Sale, product: Optional[Product] = None) -> SaleResponse:
    price = float(product.price) if product is not None else 0.0
    return SaleResponse(
        id=sale.id,
        product_id=sale.product_id,
        product_name=product.name if product is not None else None,
        product=product.name if product is not None else None,
        price=price if product is not None else None,
        revenue=round(price * int(sale.qty), 2) if product is not None else None,
        qty=int(sale.qty),
        date=sale.date,
        source=sale.source,
        created_at=sale.created_at,
    )


def _parse_excel_date(raw_val: Any) -> Optional[date]:
    if raw_val is None or raw_val == "":
        return None
    if isinstance(raw_val, datetime):
        return raw_val.date()
    if isinstance(raw_val, date):
        return raw_val
    text = str(raw_val).strip()
    if not ISO_DATE_RE.match(text):
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _parse_positive_int(raw_val: Any) -> Optional[int]:
    if raw_val is None or raw_val == "" or isinstance(raw_val, bool):
        return None
    try:
        num = float(str(raw_val).strip().replace(",", "."))
        if num <= 0 or int(round(num)) != num:
            return None
        return int(round(num))
    except (ValueError, TypeError):
        return None


class SaleUseCases:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.repo = SaleRepository(db)

    async def list_sales(
        self,
        user_id: int,
        *,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        source: Optional[str] = None,
        limit: int = 100,
        offset: int = 0,
    ) -> SaleListResponse:
        rows, total = await self.repo.list_by_user(
            user_id,
            start_date=start_date,
            end_date=end_date,
            source=source,
            limit=limit,
            offset=offset,
        )
        items = [_to_sale_response(sale, prod) for sale, prod in rows]
        return SaleListResponse(
            items=items,
            total=total,
            limit=limit,
            offset=offset,
        )

    async def create_sale(
        self,
        user: User,
        payload: SaleCreateRequest,
    ) -> SaleCreateResponse:
        if not payload.items:
            raise BadRequestException("Isi jumlah terjual minimal untuk satu menu")

        sale_date = payload.date or date.today()
        product_ids = list({item.product_id for item in payload.items})
        product_map = await self.repo.get_user_products_by_ids(user.id, product_ids)

        if len(product_map) != len(product_ids):
            raise ProductNotFoundException("Produk atau menu tidak ditemukan")

        is_pro = user.role == "pro" and (user.pro_until is None or user.pro_until >= date.today())
        used_before = await self.repo.count_monthly_sales(user.id, sale_date)
        adding_count = len(payload.items)

        if not is_pro and (used_before + adding_count) > FREE_QUOTA:
            raise FreemiumQuotaExceededException(
                f"Batas transaksi bulanan paket Freemium ({FREE_QUOTA} transaksi) telah tercapai. Silakan tingkatkan ke paket Pro."
            )

        recipes_map = await self.repo.get_recipes_for_products(product_ids)
        items_with_products = [
            (product_map[item.product_id], int(item.qty))
            for item in payload.items
        ]

        created_sales = await self.repo.create_sales_with_fefo(
            user_id=user.id,
            sale_date=sale_date,
            source=payload.source or "manual",
            items=items_with_products,
            recipes_map=recipes_map,
        )

        sale_responses = [
            _to_sale_response(s, product_map.get(s.product_id))
            for s in created_sales
        ]
        first = sale_responses[0]
        total_portions = sum(s.qty for s in sale_responses)
        used_after = used_before + len(sale_responses)

        return SaleCreateResponse(
            id=first.id,
            product_id=first.product_id,
            product_name=first.product_name,
            qty=first.qty,
            date=sale_date,
            source=payload.source or "manual",
            recorded_count=len(sale_responses),
            created=len(sale_responses),
            total=total_portions,
            quota=QuotaInfo(
                used=used_after,
                limit=None if is_pro else FREE_QUOTA,
                is_unlimited=is_pro,
            ),
            items=sale_responses,
        )

    async def _validate_and_commit_rows(
        self,
        user: User,
        raw_rows: list[tuple[int, Any, Any, Any]],
        *,
        commit: bool = True,
    ) -> SaleImportExcelResponse:
        if not raw_rows:
            raise InvalidImportDataException("Berkas Excel tidak memuat baris data transaksi")

        products = await self.repo.get_all_user_products(user.id)
        product_by_name = {p.name.strip().lower(): p for p in products}

        row_errors: list[str] = []
        error_row_numbers: list[int] = []
        valid_records: list[tuple[date, Product, int]] = []

        for row_num, raw_date, raw_menu, raw_qty in raw_rows:
            current_errors: list[str] = []

            parsed_date = _parse_excel_date(raw_date)
            if parsed_date is None:
                current_errors.append(
                    f"Baris {row_num}: Format tanggal '{raw_date}' tidak valid (harus YYYY-MM-DD)"
                )

            menu_str = str(raw_menu or "").strip()
            matched_product = product_by_name.get(menu_str.lower()) if menu_str else None
            if not menu_str:
                current_errors.append(f"Baris {row_num}: Kolom menu wajib diisi")
            elif matched_product is None:
                current_errors.append(
                    f"Baris {row_num}: Menu '{menu_str}' tidak ditemukan pada master produk"
                )

            parsed_qty = _parse_positive_int(raw_qty)
            if parsed_qty is None:
                current_errors.append(
                    f"Baris {row_num}: Jumlah porsi '{raw_qty}' tidak valid (harus bilangan bulat > 0)"
                )

            if current_errors:
                error_row_numbers.append(row_num)
                row_errors.extend(current_errors)
            elif parsed_date is not None and matched_product is not None and parsed_qty is not None:
                valid_records.append((parsed_date, matched_product, parsed_qty))

        if row_errors:
            rows_str = ", ".join(str(r) for r in error_row_numbers)
            detail_msg = "; ".join(row_errors)
            raise InvalidImportDataException(
                f"Data impor tidak valid pada baris [{rows_str}]: {detail_msg}"
            )

        if commit and valid_records:
            is_pro = user.role == "pro" and (user.pro_until is None or user.pro_until >= date.today())
            if not is_pro:
                per_month: dict[tuple[int, int], int] = {}
                for sale_dt, _, _ in valid_records:
                    key = (sale_dt.year, sale_dt.month)
                    per_month[key] = per_month.get(key, 0) + 1
                for (yr, mo), count_adding in per_month.items():
                    used = await self.repo.count_monthly_sales(user.id, date(yr, mo, 1))
                    if used + count_adding > FREE_QUOTA:
                        raise FreemiumQuotaExceededException(
                            f"Batas transaksi bulanan paket Freemium ({FREE_QUOTA} transaksi) terlampaui pada bulan {yr}-{mo:02d}."
                        )

            product_ids = list({prod.id for _, prod, _ in valid_records})
            recipes_map = await self.repo.get_recipes_for_products(product_ids)

            by_date: dict[date, list[tuple[Product, int]]] = {}
            for sale_dt, prod, qty in valid_records:
                by_date.setdefault(sale_dt, []).append((prod, qty))

            for sale_dt in sorted(by_date.keys()):
                await self.repo.create_sales_with_fefo(
                    user_id=user.id,
                    sale_date=sale_dt,
                    source="excel",
                    items=by_date[sale_dt],
                    recipes_map=recipes_map,
                )

        return SaleImportExcelResponse(
            imported_rows=len(valid_records),
            imported=len(valid_records),
            skipped=0,
            errors=[],
        )

    async def import_sales_excel(
        self,
        user: User,
        file_bytes: bytes,
    ) -> SaleImportExcelResponse:
        if not file_bytes:
            raise InvalidImportDataException("Berkas Excel kosong")

        try:
            wb = load_workbook(filename=BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise InvalidImportDataException(
                "Berkas Excel tidak dapat dibaca. Pastikan format file adalah .xlsx yang valid"
            ) from exc

        ws = wb.active
        if ws is None:
            raise InvalidImportDataException("Worksheet Excel tidak ditemukan")

        rows_iter = list(ws.iter_rows(values_only=True))
        if len(rows_iter) < 2:
            raise InvalidImportDataException("Berkas Excel tidak memiliki baris data")

        header_row = [str(cell or "").strip().lower() for cell in rows_iter[0]]
        date_idx: Optional[int] = None
        menu_idx: Optional[int] = None
        qty_idx: Optional[int] = None

        for idx, col_name in enumerate(header_row):
            if col_name in DATE_HEADER_ALIASES and date_idx is None:
                date_idx = idx
            elif col_name in MENU_HEADER_ALIASES and menu_idx is None:
                menu_idx = idx
            elif col_name in QTY_HEADER_ALIASES and qty_idx is None:
                qty_idx = idx

        if date_idx is None or menu_idx is None or qty_idx is None:
            raise InvalidImportDataException(
                "Header kolom wajib tidak lengkap. Gunakan kolom: Tanggal, Nama Menu (atau Menu), dan Jumlah Porsi Terjual (atau Jumlah)"
            )

        extracted: list[tuple[int, Any, Any, Any]] = []
        for row_num, row_vals in enumerate(rows_iter[1:], start=2):
            if not row_vals or all(cell is None or str(cell).strip() == "" for cell in row_vals):
                continue
            raw_date = row_vals[date_idx] if date_idx < len(row_vals) else None
            raw_menu = row_vals[menu_idx] if menu_idx < len(row_vals) else None
            raw_qty = row_vals[qty_idx] if qty_idx < len(row_vals) else None
            extracted.append((row_num, raw_date, raw_menu, raw_qty))

        return await self._validate_and_commit_rows(user, extracted, commit=True)

    async def import_sales_json(
        self,
        user: User,
        payload: SaleImportJsonRequest,
    ) -> SaleImportExcelResponse:
        extracted: list[tuple[int, Any, Any, Any]] = []
        for idx, row_dict in enumerate(payload.rows, start=2):
            lower_map = {str(k).strip().lower(): v for k, v in row_dict.items()}
            raw_date = next(
                (lower_map[k] for k in DATE_HEADER_ALIASES if k in lower_map),
                None,
            )
            raw_menu = next(
                (lower_map[k] for k in MENU_HEADER_ALIASES if k in lower_map),
                None,
            )
            raw_qty = next(
                (lower_map[k] for k in QTY_HEADER_ALIASES if k in lower_map),
                None,
            )
            extracted.append((idx, raw_date, raw_menu, raw_qty))

        return await self._validate_and_commit_rows(user, extracted, commit=payload.commit)

    def generate_excel_template(
        self,
        sample_menus: Optional[list[str]] = None,
    ) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Penjualan"

        headers = ["Tanggal", "Nama Menu", "Jumlah Porsi Terjual"]
        ws.append(headers)

        today = date.today()
        menus = sample_menus if sample_menus else ["Ayam Bakar Madu", "Paket Nasi Ayam Goreng"]
        for offset in (2, 1):
            sample_date = (today - timedelta(days=offset)).isoformat()
            for menu_name in menus[:3]:
                ws.append([sample_date, menu_name, 10])

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 24

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    async def get_training_history(
        self,
        user_id: int,
        *,
        history_type: str = "sales",
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        source: Optional[str] = None,
        limit: int = 500,
        offset: int = 0,
    ) -> SaleHistoryResponse:
        to_dt = end_date or date.today()
        from_dt = start_date or (to_dt - timedelta(days=13))
        events_map = await self.repo.get_events_in_range(user_id, from_dt, to_dt)

        kind = history_type.strip().lower() if history_type else "sales"
        if kind == "stock":
            batch_rows = await self.repo.list_batches_in_range(user_id, from_dt, to_dt)
            rows_out = [
                {
                    "id": b.id,
                    "date": b.received_date.isoformat(),
                    "expiry_date": b.expiry_date.isoformat(),
                    "qty": float(b.qty),
                    "unit_price": float(b.unit_price) if b.unit_price is not None else float(ing.price),
                    "source": b.source,
                    "ingredient": ing.name,
                    "unit": ing.unit,
                }
                for b, ing in batch_rows
            ]
            active_days = len({r["date"] for r in rows_out})
            return SaleHistoryResponse(
                from_date=from_dt,
                to_date=to_dt,
                start_date=from_dt,
                end_date=to_dt,
                type="stock",
                active_days=active_days,
                total_records=len(rows_out),
                events=events_map,
                rows=rows_out,
                items=[],
            )

        if kind == "waste":
            waste_rows = await self.repo.list_waste_in_range(user_id, from_dt, to_dt)
            rows_out = [
                {
                    "id": w.id,
                    "date": w.date.isoformat(),
                    "ingredient": ing.name,
                    "unit": ing.unit,
                    "qty": round(float(w.qty), 2),
                    "value": round(float(w.qty) * float(ing.price)),
                    "reason": w.reason,
                    "auto": False,
                }
                for w, ing in waste_rows
            ]
            active_days = len({r["date"] for r in rows_out})
            return SaleHistoryResponse(
                from_date=from_dt,
                to_date=to_dt,
                start_date=from_dt,
                end_date=to_dt,
                type="waste",
                active_days=active_days,
                total_records=len(rows_out),
                events=events_map,
                rows=rows_out,
                items=[],
            )

        sale_rows, total = await self.repo.list_by_user(
            user_id,
            start_date=from_dt,
            end_date=to_dt,
            source=source,
            limit=limit,
            offset=offset,
        )
        items = [_to_sale_response(sale, prod) for sale, prod in sale_rows]
        rows_out = [
            {
                "id": item.id,
                "date": item.date.isoformat(),
                "product_id": item.product_id,
                "product": item.product_name,
                "product_name": item.product_name,
                "qty": item.qty,
                "price": item.price,
                "revenue": item.revenue,
                "source": item.source,
                "created_at": item.created_at.isoformat() if item.created_at else None,
            }
            for item in items
        ]
        active_days = len({item.date for item in items})
        return SaleHistoryResponse(
            from_date=from_dt,
            to_date=to_dt,
            start_date=from_dt,
            end_date=to_dt,
            type="sales",
            active_days=active_days,
            total_records=total,
            events=events_map,
            rows=rows_out,
            items=items,
        )

    async def delete_sale(self, user_id: int, sale_id: int) -> None:
        sale = await self.repo.get_by_id_and_user(sale_id, user_id)
        if sale is None:
            raise SaleNotFoundException()
        await self.repo.delete(sale)
